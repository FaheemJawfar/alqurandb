#!/usr/bin/env python3
"""
Convert CSV translations to Excel (XLSX) format

Reads from: alqurandb_api/data/translations/csv/ (base format)
Outputs to: alqurandb_api/data/translations/xlsx/

CSV is the source of truth - this converter generates Excel files from CSV.
"""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def create_translation_excel(translation_id, csv_file_path, metadata_item, output_file):
    """Create Excel file for a single translation from CSV"""

    # Check for footnotes in CSV header
    has_footnotes = False
    with open(csv_file_path, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        if 'footnotes' in reader.fieldnames:
            has_footnotes = True

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Translation"

    # Set column widths
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 80
    if has_footnotes:
        ws.column_dimensions['D'].width = 80

    # Header row styling
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    # Add metadata at the top
    ws['A1'] = "Translation Information"
    ws['A1'].font = Font(bold=True, size=14)
    
    if has_footnotes:
        ws.merge_cells('A1:D1')
    else:
        ws.merge_cells('A1:C1')

    ws['A2'] = "ID:"
    ws['B2'] = translation_id
    ws['A3'] = "Language:"
    ws['B3'] = metadata_item.get('language', '')
    ws['A4'] = "Translator:"
    ws['B4'] = metadata_item.get('translator', '')
    if metadata_item.get('name_in_language'):
        ws['A5'] = "Native Name:"
        ws['B5'] = metadata_item.get('name_in_language', '')
        header_row = 7
    else:
        header_row = 6

    # Add verses from CSV
    verse_count = 0
    row = header_row + 1


    # Re-open to read data cleanly
    with open(csv_file_path, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        
        # Add column headers dynamically
        headers = ['Sura', 'Aya', 'Text']
        if has_footnotes:
            headers.append('Footnotes')
            
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        for data_row in reader:
            ws.cell(row=row, column=1, value=int(data_row['sura']))
            ws.cell(row=row, column=2, value=int(data_row['aya']))
            ws.cell(row=row, column=3, value=data_row['text'])
            
            if has_footnotes:
                ws.cell(row=row, column=4, value=data_row.get('footnotes', ''))

            # Align numbers to center
            ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=2).alignment = Alignment(horizontal="center")

            # Wrap text for better readability
            ws.cell(row=row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
            if has_footnotes:
                ws.cell(row=row, column=4).alignment = Alignment(wrap_text=True, vertical="top")

            row += 1
            verse_count += 1

    # Freeze the header row
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # Save workbook
    wb.save(output_file)

    return verse_count


def main():
    """Convert all CSV translations to Excel format"""

    import json

    # Paths (CSV is the base format)
    script_dir = Path(__file__).parent
    api_data_dir = script_dir.parent.parent / 'alqurandb_api' / 'data'
    metadata_file = api_data_dir / 'metadata.json'
    csv_dir = api_data_dir / 'translations' / 'csv'
    xlsx_dir = api_data_dir / 'translations' / 'xlsx'

    # Create output directory
    xlsx_dir.mkdir(parents=True, exist_ok=True)

    # Load metadata
    with open(metadata_file, 'r', encoding='utf-8') as f:
        metadata = json.load(f)

    # Create metadata lookup
    metadata_dict = {item['id']: item for item in metadata}

    # Find all CSV files
    csv_files = sorted(csv_dir.glob('*.csv'))

    if not csv_files:
        print(f"❌ No CSV files found in {csv_dir}")
        return

    print(f"Converting {len(csv_files)} translations from CSV to Excel...")

    successful = 0
    failed = 0

    for csv_file in csv_files:
        translation_id = csv_file.stem
        output_file = xlsx_dir / f"{translation_id}.xlsx"

        try:
            metadata_item = metadata_dict.get(translation_id, {})
            verse_count = create_translation_excel(translation_id, csv_file, metadata_item, output_file)

            # Get file size
            size_kb = output_file.stat().st_size / 1024

            print(f"  ✓ {translation_id}: {verse_count} verses ({size_kb:.1f} KB)")
            successful += 1

        except Exception as e:
            print(f"  ✗ {translation_id}: Error - {e}")
            failed += 1

    print(f"\n✅ Created {successful} Excel files")
    if failed > 0:
        print(f"❌ Failed: {failed}")


if __name__ == '__main__':
    main()
