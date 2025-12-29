#!/usr/bin/env python3
"""
Convert JSON translations to Excel (XLSX) format

Reads from: alqurandb_api/data/translations/json/ (base format)
Outputs to: alqurandb_api/data/translations/xlsx/

JSON is the source of truth - this converter generates Excel files from JSON.
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def create_translation_excel(translation_id, translation_data, metadata_item, output_file):
    """Create Excel file for a single translation"""

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Translation"

    # Set column widths
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 100

    # Header row styling
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    # Add metadata at the top
    ws['A1'] = "Translation Information"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:C1')

    ws['A2'] = "ID:"
    ws['B2'] = translation_id
    ws['A3'] = "Language:"
    ws['B3'] = metadata_item.get('language', '')
    ws['A4'] = "Translator:"
    ws['B4'] = metadata_item.get('translator', '')
    if metadata_item.get('name_in_language'):
        ws['A5'] = "Native Name:"
        ws['B5'] = metadata_item.get('name_in_language', '')
        header_row = 7
    else:
        header_row = 6

    # Add column headers
    headers = ['Surah', 'Ayah', 'Text']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # Add verses
    row = header_row + 1
    for key, text in translation_data.items():
        surah, ayah = key.split(':')
        ws.cell(row=row, column=1, value=int(surah))
        ws.cell(row=row, column=2, value=int(ayah))
        ws.cell(row=row, column=3, value=text)

        # Align numbers to center
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="center")

        # Wrap text for better readability
        ws.cell(row=row, column=3).alignment = Alignment(wrap_text=True, vertical="top")

        row += 1

    # Freeze the header row
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # Save workbook
    wb.save(output_file)

    return len(translation_data)


def main():
    """Create Excel files for all translations"""

    # Paths
    script_dir = Path(__file__).parent
    api_data_dir = script_dir.parent / 'alqurandb_api' / 'data'
    metadata_file = api_data_dir / 'metadata.json'
    translations_json_dir = api_data_dir / 'translations' / 'json'
    output_dir = api_data_dir / 'translations' / 'xlsx'

    # Create output directory
    output_dir.mkdir(parents=True, exist_ok=True)

    # Load metadata
    with open(metadata_file, 'r', encoding='utf-8') as f:
        metadata = json.load(f)

    print(f"Creating Excel files for {len(metadata)} translations...")

    successful = 0
    failed = 0

    for item in metadata:
        translation_id = item['id']
        json_file = translations_json_dir / f"{translation_id}.json"
        output_file = output_dir / f"{translation_id}.xlsx"

        if not json_file.exists():
            print(f"  ✗ {translation_id}: JSON file not found")
            failed += 1
            continue

        try:
            # Load translation data
            with open(json_file, 'r', encoding='utf-8') as f:
                translation_data = json.load(f)

            # Create Excel file
            verse_count = create_translation_excel(translation_id, translation_data, item, output_file)

            # Get file size
            size_kb = output_file.stat().st_size / 1024

            print(f"  ✓ {translation_id}: {verse_count} verses ({size_kb:.1f} KB)")
            successful += 1

        except Exception as e:
            print(f"  ✗ {translation_id}: Error - {e}")
            failed += 1

    print(f"\n✅ Created {successful} Excel files")
    if failed > 0:
        print(f"❌ Failed: {failed}")


if __name__ == '__main__':
    main()
